"""
Genera presentacion/TP2_PagoGlobal.pptx (16:9) con notas del orador.
Requiere:  pip install python-pptx      Ejecutar:  python3 generar_pptx.py
"""
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Inches, Pt

AQUI = Path(__file__).parent
DIAGRAMAS = AQUI.parent / "diagramas"

AZUL = RGBColor(0x0B, 0x2A, 0x5B)
AZUL_CLARO = RGBColor(0xE8, 0xF1, 0xFF)
VERDE = RGBColor(0x1F, 0x9D, 0x55)
VERDE_CLARO = RGBColor(0xE6, 0xF7, 0xEE)
NARANJA_CLARO = RGBColor(0xFF, 0xF4, 0xE0)
NARANJA = RGBColor(0xB0, 0x6A, 0x00)
GRIS = RGBColor(0x44, 0x44, 0x44)
BLANCO = RGBColor(0xFF, 0xFF, 0xFF)

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
VACIA = prs.slide_layouts[6]
TOTAL = []


def rect(slide, x, y, w, h, fill, line=None, forma=MSO_SHAPE.RECTANGLE):
    s = slide.shapes.add_shape(forma, x, y, w, h)
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
    s.shadow.inherit = False
    return s


def texto(slide, x, y, w, h, lineas, size=18, color=GRIS, bold=False, align=PP_ALIGN.LEFT,
          bullets=False):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    for i, linea in enumerate(lineas if isinstance(lineas, list) else [lineas]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(6)
        r = p.add_run()
        r.text = ("•  " + linea) if bullets else linea
        r.font.size, r.font.bold, r.font.color.rgb = Pt(size), bold, color
    return tb


def notas(slide, txt):
    slide.notes_slide.notes_text_frame.text = txt


def encabezado(slide, titulo, numero):
    rect(slide, 0, 0, prs.slide_width, Inches(1.0), AZUL)
    texto(slide, Inches(0.5), Inches(0.18), Inches(11.5), Inches(0.7), titulo, 30, BLANCO, True)
    texto(slide, Inches(11.6), Inches(0.3), Inches(1.4), Inches(0.5), f"{numero}", 16, BLANCO,
          align=PP_ALIGN.RIGHT)
    texto(slide, Inches(0.5), Inches(7.05), Inches(12), Inches(0.4),
          "TP2 · Plataformas de Pagos Electrónicos · Infraestructura global de PagoGlobal",
          11, RGBColor(0x88, 0x88, 0x88))


def slide_concepto(titulo, que_es, en_pagoglobal, tradeoff, nota):
    s = prs.slides.add_slide(VACIA)
    TOTAL.append(s)
    encabezado(s, titulo, len(TOTAL))
    # Columna izquierda: qué es
    texto(s, Inches(0.5), Inches(1.25), Inches(6), Inches(0.5), "¿Qué es?", 20, AZUL, True)
    texto(s, Inches(0.5), Inches(1.8), Inches(6.0), Inches(4.0), que_es, 20, bullets=True)
    # Columna derecha: aplicación
    rect(s, Inches(6.85), Inches(1.25), Inches(6.0), Inches(4.45), VERDE_CLARO, VERDE,
         MSO_SHAPE.ROUNDED_RECTANGLE).adjustments[0] = 0.05
    texto(s, Inches(7.1), Inches(1.35), Inches(5.6), Inches(0.5), "En PagoGlobal", 20, VERDE, True)
    texto(s, Inches(7.1), Inches(1.9), Inches(5.6), Inches(3.8), en_pagoglobal, 18, bullets=True)
    # Pie: trade-off
    rect(s, Inches(0.5), Inches(5.95), Inches(12.35), Inches(0.95), NARANJA_CLARO, NARANJA,
         MSO_SHAPE.ROUNDED_RECTANGLE).adjustments[0] = 0.15
    texto(s, Inches(0.75), Inches(6.08), Inches(11.9), Inches(0.8), "⚖  " + tradeoff, 16,
          NARANJA, True)
    notas(s, nota)


def slide_imagen(titulo, imagen, nota, alto=Inches(5.85)):
    s = prs.slides.add_slide(VACIA)
    TOTAL.append(s)
    encabezado(s, titulo, len(TOTAL))
    pic = s.shapes.add_picture(str(imagen), 0, Inches(1.1), height=alto)
    if pic.width > prs.slide_width - Inches(0.6):  # ajustar al ancho
        ratio = (prs.slide_width - Inches(0.6)) / pic.width
        pic.width, pic.height = int(pic.width * ratio), int(pic.height * ratio)
    pic.left = int((prs.slide_width - pic.width) / 2)
    notas(s, nota)


def slide_lista(titulo, items, nota, size=20):
    s = prs.slides.add_slide(VACIA)
    TOTAL.append(s)
    encabezado(s, titulo, len(TOTAL))
    texto(s, Inches(0.7), Inches(1.4), Inches(12), Inches(5.5), items, size, bullets=True)
    notas(s, nota)
    return s


def slide_tabla(titulo, filas, anchos, nota, size=14):
    s = prs.slides.add_slide(VACIA)
    TOTAL.append(s)
    encabezado(s, titulo, len(TOTAL))
    alto_fila = Inches(0.5)
    tabla = s.shapes.add_table(len(filas), len(filas[0]), Inches(0.5), Inches(1.3),
                               Inches(12.35), alto_fila * len(filas)).table
    for j, w in enumerate(anchos):
        tabla.columns[j].width = Inches(w)
    for i, fila in enumerate(filas):
        for j, valor in enumerate(fila):
            celda = tabla.cell(i, j)
            celda.text = valor
            for p in celda.text_frame.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(size)
                    r.font.bold = i == 0
                    r.font.color.rgb = BLANCO if i == 0 else GRIS
            celda.fill.solid()
            celda.fill.fore_color.rgb = AZUL if i == 0 else (AZUL_CLARO if i % 2 else BLANCO)
    notas(s, nota)


# =============================== DIAPOSITIVAS ================================
# 1. Portada
s = prs.slides.add_slide(VACIA)
TOTAL.append(s)
rect(s, 0, 0, prs.slide_width, prs.slide_height, AZUL)
texto(s, Inches(0.8), Inches(1.6), Inches(11.7), Inches(1.2), "PagoGlobal", 60, BLANCO, True)
texto(s, Inches(0.8), Inches(2.8), Inches(11.7), Inches(1.2),
      "Infraestructura para una plataforma de pagos electrónicos de alcance global", 28, BLANCO)
rect(s, Inches(0.8), Inches(4.1), Inches(3), Emu(40000), VERDE)
texto(s, Inches(0.8), Inches(4.4), Inches(11.7), Inches(1.5),
      ["Trabajo Práctico N.º 2 · Plataformas de Pagos Electrónicos",
       "Matías Giachello · Octubre 2026"], 20, RGBColor(0xC8, 0xD8, 0xF0))
notas(s, "Buenas. Voy a presentar el diseño de infraestructura de PagoGlobal, una pasarela de "
         "pagos ficticia, estilo Stripe o Mercado Pago, que opera en América, Europa y Asia. "
         "Elegí pagos porque es el dominio donde más se notan los conceptos de la materia: un "
         "cobro no se puede duplicar ni perder, y cada minuto caído es plata que no se cobra.")

slide_lista("El caso: PagoGlobal", [
    "Pasarela de pagos para comercios online: tarjetas, billeteras y transferencias",
    "Actores: comprador, comercio (API + dashboard), redes de tarjetas, agente de IA de soporte",
    "Tres regiones: San Pablo (LatAm), Virginia (Norteamérica), Irlanda (UE · GDPR)",
    "Funciones: cobrar, reembolsar, notificar al comercio, dashboard en vivo, reportes, "
    "login con Google, API para terceros",
    "Objetivo del TP: aplicar los 18 conceptos + Redis a un sitio global real",
], "Presento el caso. Lo importante es que hay tres tipos de clientes: el comprador que paga, "
   "el comercio que integra la API y mira su dashboard, y un agente de IA de soporte. Y que es "
   "global: tres regiones de AWS, con Europa separada por la ley de protección de datos.")

slide_tabla("Requisitos no funcionales", [
    ["Atributo", "Objetivo", "¿Por qué?"],
    ["Disponibilidad", "99,99 % (≈ 52 min/año)", "Estándar en pagos"],
    ["Latencia", "p99 < 300 ms", "El comprador espera en el checkout"],
    ["Throughput", "2.000 TPS normal · 20.000 TPS pico", "Black Friday / Hot Sale (×10)"],
    ["Durabilidad", "RPO = 0 en la región", "No se puede perder un cobro"],
    ["Recuperación", "RTO ≤ 15 min ante caída de región", "Conmutación a otra región"],
    ["Idempotencia", "Nunca cobrar dos veces", "Reintentos por timeouts"],
    ["Seguridad", "PCI DSS · GDPR · TLS 1.2+", "Regulación"],
], [3.0, 4.6, 4.75], "Estos números guían todas las decisiones. Fíjense en dos: el pico de 20.000 "
   "transacciones por segundo, que es lo que justifica colas y caché, y la idempotencia, que "
   "aparece en casi todos los componentes.", size=16)

slide_imagen("Arquitectura por capas", DIAGRAMAS / "arquitectura.png",
             "Este es el mapa completo. Se lee de arriba hacia abajo: clientes, edge global con "
             "DNS y CDN, la entrada de cada región con Load Balancer, API Gateway y OAuth, los "
             "microservicios en Kubernetes, la mensajería asincrónica y los datos. A la izquierda, "
             "Terraform crea todo; a la derecha, todo lo de la región se repite tres veces.")

slide_imagen("Recorrido de un pago de punta a punta", DIAGRAMAS / "flujo_pago.png",
             "Sigamos un pago. 1: el checkout baja de la CDN. 2: el POST llega con JWT y una "
             "Idempotency-Key. 3: el Gateway valida y rutea. 5-6: Redis verifica que no sea un "
             "reintento. 7: se guarda en Aurora con un UUIDv7 usando el pool, junto con el evento "
             "en la tabla outbox. 9: respondemos al comprador. Desde acá todo es asincrónico: Kafka "
             "reparte el evento a notificaciones, que manda el webhook, y a realtime, que actualiza "
             "el dashboard por WebSocket.", alto=Inches(5.0))

slide_concepto("Terraform — Infraestructura como Código", [
    "Herramienta de HashiCorp; lenguaje declarativo HCL",
    "Describís QUÉ querés; terraform plan calcula el CÓMO",
    "Guarda un estado (state) de lo que ya existe",
    "Multi-nube: AWS, GCP, Azure…",
], [
    "Un módulo \"region\" escrito una vez, instanciado 3 veces",
    "VPC, EKS, Aurora, Redis, Kafka, SQS+DLQ, S3, ALB, API Gateway, CloudFront y Route 53",
    "Cambios por pull request + plan en CI (auditoría PCI)",
    "Black Friday: eks_min_nodes = 30 y apply",
    "State cifrado en S3 con lock en DynamoDB",
], "Alternativa descartada: crear todo desde la consola → regiones distintas entre sí, "
   "sin auditoría y sin forma de reconstruir ante un desastre.",
   "Terraform es la base de un sitio global: si tengo que tener tres regiones idénticas, no las "
   "puedo hacer a mano. El código está en la carpeta terraform y lo validé con terraform validate.")

slide_concepto("CDN y Proxy inverso", [
    "CDN: red de cachés distribuida que sirve estático cerca del usuario",
    "Proxy inverso: intermediario delante del backend; oculta servidores, termina TLS, balancea",
    "La CDN es, en sí, un proxy inverso global",
], [
    "CloudFront sirve el JS/CSS del checkout: ~20 ms vs ~250 ms al origen",
    "WAF + Shield en el edge: bloquea bots, inyecciones y DDoS",
    "Proxies inversos en cadena: CloudFront → ALB → Ingress → PgBouncer",
    "La API pasa por la CDN pero sin caché (no-store)",
    "Nada de la app tiene IP pública",
], "La CDN solo cachea lo estático y público: las respuestas de pagos son privadas y dinámicas.",
   "Junté CDN y proxy inverso porque la CDN es un proxy inverso. En un Black Friday el 95 % de "
   "los bytes son estáticos: la CDN es lo que evita que se caiga el origen.")

slide_tabla("Load Balancer vs API Gateway", [
    ["", "Load Balancer", "API Gateway"],
    ["Capa", "L4 (TCP) o L7 básico", "L7 con lógica de aplicación"],
    ["Objetivo", "Repartir entre instancias IGUALES", "Rutear a servicios DISTINTOS"],
    ["Entiende de", "IPs, puertos, salud de instancias", "Rutas, JWT, cuotas, versiones"],
    ["Ejemplo", "AWS ALB / NLB, HAProxy", "AWS API Gateway, Kong"],
    ["En PagoGlobal", "ALB reparte entre réplicas del Gateway en 3 AZ",
     "/payments → Payments · /reports → Consultas · rate limit por comercio"],
], [2.4, 4.7, 5.25], "No compiten, se combinan. El Load Balancer da alta disponibilidad al "
   "Gateway, y el Gateway decide a qué microservicio va cada petición y aplica políticas como el "
   "rate limiting o la validación del JWT. Y dentro de Kubernetes hay otro nivel de balanceo "
   "entre los pods de cada servicio.", size=15)

slide_concepto("OAuth 2.0 y Session vs JWT", [
    "OAuth 2.0: autorización delegada con scopes, sin compartir la contraseña",
    "Session: cookie con ID; el estado vive en el servidor (Redis) → revocable al instante",
    "JWT: token firmado con claims; se valida sin consultar nada → escala, pero no se revoca fácil",
], [
    "Authorization Code + PKCE: una plataforma e-commerce opera en nombre del comercio",
    "Client Credentials: ERP y servicio-a-servicio",
    "OIDC: \"Iniciar sesión con Google\"",
    "API → JWT de 10 min firmado ES256, validado en cualquier región con JWKS",
    "Dashboard → Session en Redis (HttpOnly, Secure, SameSite)",
], "Usamos los dos: JWT donde importa escalar entre regiones; Session donde importa poder "
   "cerrar una cuenta robada en el acto.",
   "La pregunta típica es cuál es mejor. Mi respuesta es que depende: para la API pública uso "
   "JWT cortos, porque cualquier región los valida con la clave pública. Para el dashboard uso "
   "sesión, porque si a un comercio le roban la cuenta tengo que poder cortarla ya.")

slide_concepto("API, Webhook y MCP", [
    "API: solicitud-respuesta; el cliente pregunta",
    "Webhook: \"API inversa\"; el servidor avisa con un POST cuando pasa algo",
    "MCP: protocolo abierto (Anthropic, 2024) para que un LLM use herramientas y datos de forma "
    "estandarizada",
], [
    "API REST /v1 para crear pagos y reembolsos",
    "Webhook payment.approved firmado con HMAC-SHA256 + timestamp (anti-replay)",
    "Entrega at-least-once → el comercio deduplica por event_id",
    "Servidor MCP con buscar_pago y motivo_rechazo: solo lectura, token con scope payments:read",
], "Sin webhooks el comercio haría polling: miles de consultas inútiles. "
   "Al LLM nunca se le da una herramienta que mueva dinero.",
   "Tres formas de integrarse. En el código hay un ejemplo que muestra que si un atacante "
   "modifica el monto del webhook, la firma no coincide y se rechaza con 401. Y el servidor MCP "
   "permite que un asistente responda por qué se rechazó un pago con datos reales.")

slide_concepto("WebSocket vs Webhook", [
    "WebSocket: conexión persistente, bidireccional, full-duplex sobre TCP",
    "Webhook: un POST por evento hacia el servidor de un tercero",
    "WS → a un navegador conectado ahora · Webhook → a un backend",
], [
    "WebSocket: el dashboard muestra cada venta al instante",
    "Webhook: el ERP del comercio marca el pedido como pagado",
    "Los dos salen del MISMO evento en Kafka (dos consumidores)",
    "Varios pods de Realtime: Redis Pub/Sub entre pods",
    "ALB con timeout largo para conexiones WS",
], "Si el receptor está caído: el webhook se reintenta (y DLQ); el WS se pierde, por eso "
   "al reconectar el dashboard pide lo último por API.",
   "Parecen parecidos pero tienen destinatarios distintos: WebSocket es para una persona que "
   "está mirando la pantalla; webhook es para un sistema.")

slide_concepto("Kafka y arquitectura asincrónica con colas", [
    "Kafka: log distribuido, particionado y replicado; los mensajes se retienen y se pueden releer",
    "Cola (SQS/RabbitMQ): tareas con reintento por mensaje",
    "Asincronía = desacople + tolerancia a fallos + absorber picos",
], [
    "Regla: lo que el comprador necesita para ver \"aprobado\" es sincrónico; el resto, asincrónico",
    "Tópico payments.events, partición = merchant_id (orden por comercio)",
    "5 grupos de consumo: notificaciones, consultas, realtime, antifraude, contabilidad",
    "Outbox: pago y evento en la MISMA transacción",
    "acks=all, idempotencia del productor, 3 réplicas",
], "Todo sincrónico sería más simple, pero si el servicio de emails se cae, se caerían los cobros.",
   "Esta es la decisión que más impacta. Agregar un nuevo consumidor, por ejemplo un programa "
   "de puntos, no requiere tocar Payments. Y el patrón outbox evita el problema de guardar el "
   "pago pero perder el evento.")

slide_concepto("Dead Letter Queue", [
    "Cola secundaria para mensajes que fallaron N veces",
    "Evita reintentos infinitos y que un mensaje \"veneno\" bloquee la cola",
    "Nada se pierde: queda para análisis y reproceso (redrive)",
], [
    "SQS webhooks con maxReceiveCount = 8 (~24 h con backoff)",
    "Error transitorio (timeout, 503) → reintento",
    "Error permanente (404, payload inválido) → directo a la DLQ",
    "Alarma si la DLQ tiene mensajes; el comercio puede reenviarlos desde el dashboard",
    "En Kafka: tópico payments.events.dlq",
], "Sin DLQ: o descartás el mensaje (perdés la notificación de un pago) o reintentás para siempre.",
   "En el ejemplo de código consumidor_dlq.py se ve en vivo: un comercio con la URL mal va directo "
   "a la DLQ y uno caído agota los reintentos y termina ahí, sin frenar a los demás.")

slide_concepto("CQRS — separar escrituras de lecturas", [
    "Command: cambia estado · Query: solo lee",
    "Cada lado con su modelo y su base optimizada",
    "Consecuencia: consistencia eventual en las lecturas",
], [
    "Escritura: Aurora PostgreSQL normalizada (ledger)",
    "Lectura: OpenSearch con documentos listos para el dashboard",
    "Se sincroniza consumiendo payments.events",
    "Reportes de 90 días no frenan a los pagos",
    "El estado puntual de un pago se lee del lado de comandos",
], "Trade-off: un pago tarda ~1 s en aparecer en reportes. Aceptable para reportes, no para el "
   "estado del cobro.",
   "Sin CQRS, cuando a fin de mes todos los comercios bajan reportes, esas consultas pesadas "
   "competirían con los cobros en la misma base.")

slide_concepto("Redis — la memoria rápida del sistema", [
    "Base clave-valor en RAM, latencia < 1 ms",
    "Estructuras: strings, hashes, listas, sets, sorted sets, streams",
    "TTL por clave; comandos atómicos (núcleo single-thread)",
    "Persistencia opcional (RDB / AOF)",
], [
    "Idempotencia: SET idem:{key} … NX EX 86400",
    "Rate limiting: INCR + EXPIRE por comercio y segundo",
    "Sesiones del dashboard (revocación con DEL)",
    "Caché cache-aside de configuración del comercio",
    "Pub/Sub entre pods de WebSocket",
], "Redis NO es la fuente de verdad: si se cae, el sistema va a la base (más lento, pero sin "
   "perder pagos).",
   "Redis aparece en cinco lugares. El más importante para pagos es la idempotencia: con SET NX "
   "solo la primera petición gana; si el cliente reintenta, recibe la misma respuesta y no se "
   "cobra dos veces. En la demo, la primera consulta tarda 65 ms y la segunda menos de 1 ms.")

slide_concepto("Connection Pool", [
    "Conjunto de conexiones abiertas y reutilizables",
    "Abrir una conexión a PostgreSQL = TCP + TLS + auth + proceso nuevo",
    "Préstamo y devolución; si se agota, se espera con timeout",
], [
    "Black Friday: 200 pods × 20 conexiones = 4.000 → PostgreSQL colapsa",
    "Pool local de 10 por pod + PgBouncer (transaction pooling)",
    "≈ 300 conexiones reales a Aurora",
    "Timeout corto: si no hay conexión, 503 rápido",
    "Se cierra al final del Graceful Shutdown",
], "Pool demasiado grande: satura la base. Demasiado chico: las peticiones esperan. Se mide y se ajusta.",
   "El problema no es una sola app, es que cuando escalamos horizontalmente se multiplican las "
   "conexiones. Por eso hay dos niveles de pool.")

slide_tabla("Secuencia vs UUID (y su relación con el Load Balancer)", [
    ["", "Secuencia", "UUIDv4", "UUIDv7 ✔"],
    ["¿Quién lo genera?", "La base (central)", "Cualquier instancia", "Cualquier instancia"],
    ["¿Adivinable?", "Sí (IDOR)", "No", "No"],
    ["Revela volumen", "Sí", "No", "No"],
    ["Índice B-tree", "Excelente", "Malo (aleatorio)", "Muy bueno (ordenado por tiempo)"],
    ["Multi-región", "Conflictos", "Sin conflictos", "Sin conflictos"],
    ["LB + IDs", "Todas las instancias dependen de la base",
     "Cada instancia genera sola", "Ídem + Idempotency-Key compartida en Redis"],
], [2.6, 3.1, 3.0, 3.65], "Elegí UUIDv7, del RFC 9562. Es imposible de enumerar, se genera en "
   "cualquier instancia sin ir a la base, y como empieza con el timestamp no rompe el índice. "
   "La relación con el Load Balancer: como el LB manda cada petición a cualquier instancia, "
   "conviene que cualquier instancia pueda generar IDs sin coordinar. La secuencia la dejo solo "
   "para el número de factura, que por ley debe ser correlativo.", size=14)

slide_concepto("Presigned URL", [
    "URL de S3 con firma criptográfica y vencimiento",
    "Permite UNA acción sobre UN objeto durante poco tiempo",
    "Sin credenciales para el cliente y sin bucket público",
], [
    "Descarga del reporte de liquidación: URL válida 5 min",
    "Subida de evidencia de contracargo: POST prefirmado, máx. 10 MB, solo PDF",
    "El archivo va directo navegador ↔ S3: no consume CPU de los servicios",
    "Bucket privado, cifrado con KMS, auditado en CloudTrail",
], "Alternativa descartada: pasar los archivos por la app (cara y lenta) o hacer público el "
   "bucket (inaceptable con datos de comercios).",
   "El servicio verifica permisos y firma la URL. Generar la URL es local, no hace ninguna "
   "llamada de red.")

slide_concepto("Graceful Shutdown", [
    "Al recibir SIGTERM: dejar de aceptar → terminar lo que está en curso → cerrar recursos → salir",
    "Kubernetes lo envía en cada despliegue o al achicar el cluster",
    "terminationGracePeriodSeconds + preStop + readiness probe",
], [
    "/ready pasa a 503 → el ALB saca al pod de rotación",
    "Peticiones nuevas: 503 \"reintente\" (van a otro pod)",
    "El pago en curso termina y responde 201",
    "Commit de offsets de Kafka, cierre del pool, WS con código 1001",
    "Deregistration delay de 30 s en el target group",
], "Sin esto: un pod muere entre la aprobación del banco y el INSERT → comprador cobrado sin registro.",
   "En la demo mando kill -TERM mientras hay un pago en curso: el pago responde approved, las "
   "peticiones nuevas reciben 503 y recién después se cierra el pool y el proceso termina.")

slide_lista("Estrategia global y recuperación ante desastres", [
    "Route 53 por latencia: cada cliente va a la región sana más cercana",
    "Región de origen por comercio: sus pagos se escriben en un solo lugar (ledger consistente)",
    "Residencia de datos: comercios UE en eu-west-1, sin replicar fuera (GDPR)",
    "Cae un pod / nodo / zona → réplicas en 3 AZ, impacto de segundos",
    "Cae una región → Aurora Global DB (< 1 s de retraso) + MSK Replicator → RTO ≤ 15 min",
    "PCI DSS: tokenización del PAN en un vault aislado, HSM, mTLS, MFA, secretos rotados",
    "Simulacros trimestrales: un plan de DR no probado no existe",
], "Lo difícil de un sitio de pagos global es la consistencia del ledger. Por eso cada comercio "
   "tiene una región de origen: se lee desde cualquier lado, pero se escribe en un solo lugar.",
   size=19)

slide_tabla("Prueba de fuego: Black Friday (×10)", [
    ["Problema", "Qué lo resuelve"],
    ["Avalancha de visitas al checkout", "CDN"],
    ["Bots y abuso", "WAF + rate limiting en el API Gateway con Redis"],
    ["Reintentos por lentitud → cobro doble", "Idempotency-Key en Redis"],
    ["Miles de escrituras simultáneas", "Connection Pool + PgBouncer"],
    ["Comercios refrescando el dashboard", "CQRS + WebSocket (en vez de polling)"],
    ["Emails, webhooks, comisiones", "Kafka y colas absorben el pico"],
    ["Servidor de un comercio caído", "Reintentos + DLQ"],
    ["Autoscaling sube y baja pods", "Graceful Shutdown"],
    ["Pre-escalar la infraestructura", "Terraform (eks_min_nodes = 30)"],
], [6.0, 6.35], "Esta tabla resume por qué cada pieza está ahí: frente a un pico de diez veces el "
   "tráfico normal, cada concepto resuelve un problema concreto.", size=15)

slide_lista("Conclusiones", [
    "Ningún concepto alcanza solo: la robustez sale de combinarlos",
    "En pagos mandan la idempotencia y la consistencia: ledger transaccional, UUIDv7, outbox, "
    "consumidores idempotentes",
    "Separar lo sincrónico de lo asincrónico es la decisión de mayor impacto",
    "Cada decisión es un trade-off: CQRS → consistencia eventual; JWT → difícil revocar; "
    "UUID → más espacio",
    "Un sitio global no se construye a mano: Terraform hace las regiones idénticas y reconstruibles",
], "Cierro con esto: lo que aprendí es que no hay una tecnología correcta, sino decisiones con "
   "costos, y lo importante es elegir conscientemente y poder justificarlas.")

s = prs.slides.add_slide(VACIA)
TOTAL.append(s)
rect(s, 0, 0, prs.slide_width, prs.slide_height, AZUL)
texto(s, Inches(0.8), Inches(2.6), Inches(11.7), Inches(1.4), "¿Preguntas?", 60, BLANCO, True,
      PP_ALIGN.CENTER)
texto(s, Inches(0.8), Inches(4.1), Inches(11.7), Inches(1.0),
      "Código Terraform, ejemplos en Python y diagramas en la carpeta tp2-infraestructura-global",
      18, RGBColor(0xC8, 0xD8, 0xF0), align=PP_ALIGN.CENTER)
notas(s, "Gracias. ¿Preguntas?")

salida = AQUI / "TP2_PagoGlobal.pptx"
prs.save(salida)
print(f"{salida.name}: {len(TOTAL)} diapositivas")
